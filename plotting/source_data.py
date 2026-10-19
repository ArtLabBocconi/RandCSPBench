"""Build the Source Data workbook of the paper: one sheet per figure, one row per plotted point.

The values are recomputed from the per-instance results in results_evaluation/ with the same
selections, aggregations and fits used in paper_figures.ipynb.

    uv run python plotting/source_data.py  # writes plotting/Source_Data.xlsx
"""
import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Font
from scipy.optimize import curve_fit

RESULTS = Path(__file__).resolve().parent.parent / "results_evaluation"
REPO_URL = "https://github.com/ArtLabBocconi/RandCSPBench"
PAPER_TITLE = "Benchmarking Graph Neural Networks in Solving Hard Constraint Satisfaction Problems"

SAT_ALGOS = {
    "FMS_nsteps=100N^2": "FMS",
    "NeuroSAT_unsupervised_niters=2N_Ntrain=256": "NeuroSAT",
    "QuerySAT_niters=2N_Ntrain=256": "QuerySAT",
    "SimulatedAnnealing_nsteps=1000N^2": "SA",
    "SurveyPropagation_frac-decimated=0.00125": "SP",
}
COL_ALGOS = {
    "BeliefPropagation_frac-decimated=1overN": "BP",
    "FMScol": "FMS",
    "QuerySAT_nsteps=2N": "QuerySAT",
    "SimulatedAnnealing_nsteps=1000N^2": "SA",
    "rPI-GNN_ntrials=5": "rPI-GNN",
}
FMS_COL = {"3COL": "FMScol_eta=0.37_nsteps=625N^2", "5COL": "FMScol_eta=0.25_nsteps=625N^2"}
PROBLEM_NAMES = {"3SAT": "3-SAT", "4SAT": "4-SAT", "3COL": "3-col", "5COL": "5-col"}

X = "alpha / c"
P_COLS = {"P": "P", "P_err": "P error"}
E_COLS = {"E_mean": "E", "E_err": "E error"}


def load(problem, algo):
    if algo == "FMScol":
        algo = FMS_COL[problem]
    df = pd.read_csv(RESULTS / problem / f"{problem}_{algo}.csv")
    # M counts edges for coloring, hence c = 2M/N
    df[X] = df.M / df.N * (2 if problem.endswith("COL") else 1)
    return df


def aggregate(df):
    df = df.assign(solved=(df.E == 0).astype(int))
    sem = lambda s: s.std() / np.sqrt(len(s))
    return df.groupby(["N", "M"]).agg(
        **{X: (X, "first")},
        Instances=("E", "size"),
        Solved=("solved", "sum"),
        P=("solved", "mean"),
        P_err=("solved", sem),
        E_mean=("E", "mean"),
        E_err=("E", sem),
    ).reset_index()


def finalize(df, value_cols):
    keys = [c for c in df.columns if c not in ("Instances", "Solved", *P_COLS, *E_COLS)]
    counts = ["Instances", "Solved"] if "P" in value_cols else ["Instances"]
    return df[keys + counts + list(value_cols)].rename(columns=value_cols)


def fixed_size_panels():
    panels = zip("abcdefgh", [(p, n) for p in ("3SAT", "4SAT", "3COL", "5COL") for n in (256, 1024)])
    frames = []
    for panel, (problem, N) in panels:
        algos = SAT_ALGOS if problem.endswith("SAT") else COL_ALGOS
        for algo, name in algos.items():
            df = load(problem, algo)
            df = df[df.N == N]
            if problem == "3COL":
                df = df[df[X] > 3.25]
            agg = aggregate(df)
            agg.insert(0, "Algorithm", name)
            agg.insert(0, "Problem", PROBLEM_NAMES[problem])
            agg.insert(0, "Panel", panel)
            frames.append(agg)
    return pd.concat(frames, ignore_index=True)


def fig1():
    frames = []
    for panel, (training, niters) in zip(
        "abcdef", [(t, i) for t in ("supervised", "unsupervised") for i in ("32", "512", "2N")]
    ):
        df = load("3SAT", f"NeuroSAT_{training}_niters={niters}_Ntrain=256")
        agg = aggregate(df[df.N <= 256])
        agg.insert(0, "Iterations", niters)
        agg.insert(0, "Training", training)
        agg.insert(0, "Panel", panel)
        frames.append(agg)
    return [(None, finalize(pd.concat(frames, ignore_index=True), P_COLS))]


def fig2():
    return [(None, finalize(fixed_size_panels(), P_COLS))]


def supp_fig1():
    df = fixed_size_panels()
    # SP returns no configuration when it does not converge, so it has no energy curve
    return [(None, finalize(df[df.Algorithm != "SP"], E_COLS))]


def logistic(x, Q, b, v, x_N):
    return 1 / ((1 + Q * np.exp(-b * v * (x - x_N))) ** (1 / v))


def fig3():
    panels = [
        ("a", "3SAT", "QuerySAT_niters=2N_Ntrain=256", "QuerySAT", 16, (2.0, 6.0), 4.1),
        ("b", "4SAT", "SurveyPropagation_frac-decimated=0.00125", "SP", 128, (7.0, 11.0), 9.7),
        ("c", "3COL", "rPI-GNN_ntrials=5", "rPI-GNN", 128, (2.0, 5.0), None),
    ]
    points, fits = [], []
    for panel, problem, algo, name, Nmin, (lo, hi), threshold in panels:
        df = load(problem, algo)
        agg = aggregate(df[df.N >= Nmin])
        for N, g in agg.groupby("N"):
            (Q, b, v, x_N), _ = curve_fit(
                logistic, g[X].values, g.P.values,
                bounds=([0.0, -np.inf, 0.0, lo], [np.inf, np.inf, np.inf, hi]), maxfev=10000,
            )
            fits.append(dict(Panel=panel, Problem=PROBLEM_NAMES[problem], Algorithm=name, N=N,
                             Q=Q, b=b, v=v, x_N=x_N, **{"Algorithmic threshold (dashed line)": threshold}))
        agg.insert(0, "Algorithm", name)
        agg.insert(0, "Problem", PROBLEM_NAMES[problem])
        agg.insert(0, "Panel", panel)
        points.append(agg)
    return [
        ("Data points", finalize(pd.concat(points, ignore_index=True), P_COLS)),
        ("Fitted curves: P(x) = [1 + Q exp(-b v (x - x_N))]^(-1/v), x = alpha or c", pd.DataFrame(fits)),
    ]


def supp_fig2():
    panels = [("a, b", "3SAT", "SurveyPropagation_frac-decimated=0.00125", "SP"),
              ("c, d", "4SAT", "SurveyPropagation_frac-decimated=0.00125", "SP"),
              ("e, f", "3COL", "FMScol", "FMS"),
              ("g, h", "5COL", "FMScol", "FMS")]
    frames = []
    for panel, problem, algo, name in panels:
        df = load(problem, algo)
        agg = aggregate(df[df.N >= 4096])
        agg.insert(0, "Algorithm", name)
        agg.insert(0, "Problem", PROBLEM_NAMES[problem])
        agg.insert(0, "Panels (P, E)", panel)
        frames.append(agg)
    return [(None, finalize(pd.concat(frames, ignore_index=True), P_COLS | E_COLS))]


def supp_fig8():
    series = [
        ("left", 20, 2, 25600, "hiddim=20_nlayers=2"),
        ("left", 30, 2, 25600, "hiddim=30_nlayers=2"),
        ("left", 40, 2, 25600, "hiddim=40_nlayers=2"),
        ("left", 50, 2, 25600, "hiddim=50_nlayers=2"),
        ("right", 50, 2, 25600, "hiddim=50_nlayers=2"),
        ("right", 50, 3, 25600, "hiddim=50_nlayers=3"),
        ("right", 50, 2, 100000, "hiddim=50_nlayers=2_max_nepochs"),
    ]
    frames = []
    for panel, hiddim, nlayers, nep_max, suffix in series:
        df = load("3COL", f"rPI-GNN_ntrials=5_{suffix}")
        agg = aggregate(df[df.N == 256])
        agg.insert(0, "Max epochs (nep_max)", nep_max)
        agg.insert(0, "Hidden layers (kappa)", nlayers)
        agg.insert(0, "Hidden dimension (d_h)", hiddim)
        agg.insert(0, "Panel", panel)
        frames.append(agg)
    return [(None, finalize(pd.concat(frames, ignore_index=True), P_COLS))]


ERR = "error bars: standard deviation over instances divided by sqrt(Instances), i.e. the standard error of the mean"
SHEETS = [
    ("Fig. 1", fig1,
     "Probability P of finding a satisfying assignment for 3-SAT with NeuroSAT, trained supervised (a-c) "
     f"or unsupervised (d-f), run for 32, 512 or 2N message-passing iterations; {ERR}."),
    ("Fig. 2", fig2,
     "Probability P of finding a solution vs alpha (K-SAT) or c (q-col) at N = 256 and N = 1024 for each "
     f"algorithm; {ERR}."),
    ("Fig. 3", fig3,
     "Algorithmic threshold estimation: P vs alpha or c for several N, with the generalized logistic fits drawn "
     f"as curves; {ERR}."),
    ("Supplementary Fig. 1", supp_fig1,
     "Average extensive energy E (number of violated constraints of the returned configuration) at N = 256 and "
     f"N = 1024 for each algorithm; {ERR}."),
    ("Supplementary Fig. 2", supp_fig2,
     "Solving probability P and average extensive energy E of SP (K-SAT) and FMS (q-col) for N >= 4096; "
     f"{ERR}."),
    ("Supplementary Fig. 8", supp_fig8,
     "rPI-GNN on 3-col at N = 256: P vs c when varying the hidden dimension d_h (left) or the number of hidden "
     f"layers kappa and the maximum number of epochs (right); {ERR}."),
]

README = [
    ("Source Data", None),
    (f"Paper: {PAPER_TITLE} (Skenderi et al.)", None),
    ("Each sheet holds the values plotted in one figure, one row per data point; the 'Panel' column refers to "
     "the panel labels in the figure.", None),
    ("Columns", "bold"),
    ("N: number of variables (K-SAT) or nodes (q-col). M: number of clauses (K-SAT) or edges (q-col).", None),
    ("alpha / c: control parameter, alpha = M/N for K-SAT and mean connectivity c = 2M/N for q-col.", None),
    ("Instances: number of test instances at that point. Solved: instances where the algorithm returned a "
     "zero-energy configuration.", None),
    ("P: Solved / Instances. E: number of violated constraints of the returned configuration, averaged over "
     "instances.", None),
    ("P error, E error: standard deviation over instances divided by sqrt(Instances), as drawn in the error "
     "bars.", None),
    ("Raw data", "bold"),
    (f"Per-instance results (one row per instance with its final energy) and the code that produced this file "
     f"(plotting/source_data.py) are available at {REPO_URL}.", None),
]


def write_block(ws, writer, sheet, row, title, df):
    if title:
        ws.cell(row=row, column=1, value=title).font = Font(bold=True)
        row += 1
    df.to_excel(writer, sheet_name=sheet, startrow=row - 1, index=False)
    return row + len(df) + 2


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--output", default=Path(__file__).resolve().parent / "Source_Data.xlsx")
    args = parser.parse_args()

    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        pd.DataFrame().to_excel(writer, sheet_name="README", index=False)
        ws = writer.sheets["README"]
        for i, (text, style) in enumerate(README, start=1):
            ws.cell(row=i, column=1, value=text).font = Font(bold=style == "bold" or i == 1, size=14 if i == 1 else 11)
        ws.column_dimensions["A"].width = 140

        for sheet, build, description in SHEETS:
            blocks = build()
            pd.DataFrame().to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            ws.cell(row=1, column=1, value=sheet).font = Font(bold=True, size=14)
            ws.cell(row=2, column=1, value=description)
            row = 4
            for title, df in blocks:
                row = write_block(ws, writer, sheet, row, title, df)
            for col in ws.iter_cols(min_row=4):
                ws.column_dimensions[col[0].column_letter].width = 14
            print(f"{sheet}: {sum(len(df) for _, df in blocks)} rows")


if __name__ == "__main__":
    main()
